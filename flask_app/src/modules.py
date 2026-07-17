from dataclasses import dataclass
import os
import pandas as pd
from itertools import groupby
from pptx import Presentation
from watchdog.observers import Observer
from watchdog.events import FileSystemEventHandler
from datetime import date
import threading
from utils import debounce
import logging
from utils import check_not_blank, get_urls_in_str, check_int
from config import get_conference_and_modules_path

@dataclass
class ModuleTag:
    category: str
    tag: str

@dataclass
class TagCategory:
    name: str
    tags: list[ModuleTag]

@dataclass
class ModuleMessage:
    slide_index1: int
    content: str

@dataclass
class ModuleSource:
    slide_index1: int
    title: str
    value: str | None
    source: str
    urls: list[str]
    expiry_date: date | None

@dataclass
class ModuleChange:
    slide_index1: int
    description: str
    date: date | None

@dataclass
class ModuleMaintainerInfo:
    name: str
    contact_email: str | None
    contact_discord: str | None

@dataclass
class ConferenceModule:
    id: int
    title: str
    description: str
    duration_minutes: int
    img_url: str
    slides_path: str
    tags: list[ModuleTag]
    has_cover_slide: bool
    messages: list[ModuleMessage]
    sources: list[ModuleSource]
    changes: list[ModuleChange]
    slides_count: int
    maintainer: ModuleMaintainerInfo
    is_valid: bool
    invalid_reason: str | None
    is_hidden: bool

class ConferenceModulesService:
    modules: list[ConferenceModule]
    conferences_and_modules_path: str

    def __init__(self, conferences_and_modules_path: str):
        self.modules = read_modules(conferences_and_modules_path)
        self.conferences_and_modules_path = conferences_and_modules_path
        threading.Thread(target=lambda: self._udpate_modules_regularly()).start()
        

    def get_modules(self) -> list[ConferenceModule]: 
        return self.modules
    
    def get_valid_modules(self) -> list[ConferenceModule]: 
        return list(filter(lambda mod: mod.is_valid, self.modules))
    
    def get_visible_modules(self) -> list[ConferenceModule]: 
        return list(filter(lambda mod: not mod.is_hidden, self.modules))
    
    def get_module_by_id(self, module_id: int) -> ConferenceModule:
        module = next(m for m in self.get_modules() if m.id == module_id)
        if module is None:
            raise ValueError(f'module with id {module_id} not found')
        return module
    
    def are_all_modules_valid(self, module_ids: list[int]) -> bool:
        for module_id in module_ids:
            try:
                module = self.get_module_by_id(module_id)
                if not module.is_valid:
                    return False
            except ValueError:
                return False
        return True
    
    def _udpate_modules_regularly(self):

        @debounce(5)
        def debouced_modules_update():
            logging.info('File change detected, updating modules...')
            self.modules = read_modules(self.conferences_and_modules_path)
            logging.info('Modules updated !')

        class UpdateModulesHandler(FileSystemEventHandler):
            def on_any_event(self, _):
                debouced_modules_update()

        path_to_watch = os.path.join(self.conferences_and_modules_path, 'modules')

        observer = Observer()
        observer.schedule(UpdateModulesHandler(), path_to_watch, recursive=True)
        observer.start()

        print(f'Watching folder `{path_to_watch}` for modules updates...')

        try:
            while observer.is_alive():
                observer.join(1)
        finally:
            observer.stop()
            observer.join()

def read_modules(folder) -> list[ConferenceModule]:
    modules_folder = os.path.join(folder, 'modules')
    modules_subfolders = os.listdir(modules_folder)
    modules: list[ConferenceModule] = []
    for module_subfolder in modules_subfolders:
        try:
            module_path = os.path.join(modules_folder, module_subfolder)
            
            module_files = [f for f in os.listdir(module_path) if os.path.isfile(os.path.join(module_path, f))]

            module_definition_file = next((x for x in module_files if x.endswith('module.xlsx')), None)
            if (module_definition_file is None):
                raise ValueError(f'Unable to find a file ending with `.module.xlsx` in module `{module_subfolder}`')
            module_definition_file_path = os.path.join(module_path, module_definition_file)

            module_slides_file = next((x for x in module_files if x.endswith('slides.pptx')), None)
            if (module_slides_file is None):
                raise ValueError(f'Unable to find a file ending with `.slides.pptx` in module `{module_subfolder}`')
            module_slides_file_path = os.path.join(module_path, module_slides_file)

            prs = Presentation(module_slides_file_path)
            slides_count = len(prs.slides)

            pd_xl_file = pd.ExcelFile(module_definition_file_path)
            general_df = pd.read_excel(pd_xl_file, 'General', header=None)
            tags_df = pd.read_excel(pd_xl_file, 'Etiquettes')
            messages_df = pd.read_excel(pd_xl_file, 'Messages clés pour évaluation')
            sources_df = pd.read_excel(pd_xl_file, 'Données & sources')
            maintainer_df = pd.read_excel(pd_xl_file, 'Equipe', header=None)
            changes_df = pd.read_excel(pd_xl_file, 'Changelog')

            module_cover_file = next((x for x in module_files if x.endswith('cover.png') or x.endswith('cover.jpg')), None)
            if (module_cover_file is None):
                raise ValueError(f'Unable to find a cover file ending with `.cover.png` or `.cover.jpg` in module `{module_subfolder}`')

            pd_xl_file.close()

            check_not_blank(general_df[1][4], "Statut")

            modules.append(ConferenceModule(
                id=int(module_subfolder[0:4]),
                title=check_not_blank(general_df[1][0], "Titre"),
                description=check_not_blank(general_df[1][1], "Description"),
                duration_minutes=check_int(general_df[1][2], "Durée (minutes)"),
                img_url=f'/static/modules/{module_subfolder}/{module_cover_file}',
                slides_path=f'{modules_folder}/{module_subfolder}/{module_slides_file}',
                tags=[ModuleTag(category, tag) for category in tags_df.columns for tag in tags_df[category].tolist() if isinstance(tag, str)],
                has_cover_slide=general_df[1][3],
                messages=[ModuleMessage(row.iloc[0], row.iloc[1]) for _, row in messages_df.iterrows()],
                sources=[ModuleSource(
                    slide_index1=row.iloc[0], 
                    title=row.iloc[1], 
                    value=row.iloc[2] if not pd.isna(row.iloc[2]) else None, 
                    source=row.iloc[3], 
                    urls=get_urls_in_str(row.iloc[3]) if row.iloc[3] else [],
                    expiry_date=row.iloc[4].date() if not pd.isna(row.iloc[4]) else None
                ) for _, row in sources_df.iterrows()],
                changes=[ModuleChange(
                    slide_index1=row.iloc[0], 
                    description=row.iloc[1],
                    date=row.iloc[2].date() if not pd.isna(row.iloc[2]) else None
                ) for _, row in changes_df.iterrows()],
                slides_count=slides_count,
                maintainer=ModuleMaintainerInfo(
                    name=maintainer_df[0][1],
                    contact_email=maintainer_df[1][1],
                    contact_discord=maintainer_df[2][1],
                ),
                is_valid=True,
                invalid_reason=None,
                is_hidden=(general_df[1][4] == 'Caché')
            ))
        except Exception as e:
            try:
                module_id =int(module_subfolder[0:4])
                modules.append(ConferenceModule(
                    id=module_id,
                    title='',
                    description='',
                    duration_minutes=0,
                    img_url='',
                    slides_path='',
                    tags=[],
                    has_cover_slide=False,
                    messages=[],
                    sources=[],
                    changes=[],
                    slides_count=0,
                    maintainer=ModuleMaintainerInfo(
                        name='Inconnu',
                        contact_email=None,
                        contact_discord=None,
                    ),
                    is_valid=False,
                    invalid_reason=str(e),
                    is_hidden=True
                ))
                logging.warning(f'An error occured while parsing module {module_id}: {str(e)}')
            except Exception as e2:
                logging.error(f'Completely unable to parse module from folder: `{module_subfolder}`: {str(e)}, {str(e2)}')


    return modules

def get_all_tags(modules: list[ConferenceModule]) -> list[TagCategory]:
    tags: list[ModuleTag] = []
    for module in modules:
        tags = tags + module.tags
    return  [TagCategory(name=k, tags=unique_by_key(v, lambda t: t.tag)) for k, v in groupby(sorted(tags, key=lambda m:m.category), lambda m:m.category)]


def unique_by_key(list, getkey):
    seen = set()
    return [seen.add(getkey(obj)) or obj for obj in list if getkey(obj) not in seen]


conference_modules_service = ConferenceModulesService(get_conference_and_modules_path())